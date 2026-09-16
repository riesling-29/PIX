"""Strict table readers with evidence from the source being parsed.

Delimited files retain text cells. Excel and record inputs retain their native
cell types; deciding what those cells mean belongs to the mapping layer.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import math
import os
import posixpath
import zipfile
import zlib
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from pix.ocel.ingest.contract import ImportStage, ImportStatus, Transformation

_MAX_XLSX_CELLS = 10_000_000


@dataclass(frozen=True, slots=True)
class TableData:
    """A materialized table and the identity of its source representation."""

    fields: tuple[str, ...]
    rows: tuple[dict[str, object], ...]
    source: str
    kind: str
    sha256: str
    size: int
    transformations: tuple[Transformation, ...] = ()


class TableSourceError(ValueError):
    """A source failure carrying every available source evidence field."""

    def __init__(
        self,
        message: str,
        *,
        status: ImportStatus,
        stage: ImportStage,
        code: str,
        at: tuple[str, ...] = (),
        source: str = "memory:records",
        kind: str | None = "records",
        sha256: str | None = None,
        size: int | None = None,
    ) -> None:
        super().__init__(message)
        self.status = status
        self.stage = stage
        self.code = code
        self.at = at
        self.source = source
        self.kind = kind
        self.sha256 = sha256
        self.size = size


def _invalid(
    message: str,
    code: str,
    at: tuple[str, ...] = (),
    *,
    status: ImportStatus = ImportStatus.SCHEMA_INVALID,
    stage: ImportStage = ImportStage.SCHEMA,
) -> TableSourceError:
    return TableSourceError(message, status=status, stage=stage, code=code, at=at)


def load_table(
    source: object,
    *,
    format: str | None = None,
    sheet: str | int | None = None,
    delimiter: str = ",",
    encoding: str = "utf-8",
) -> TableData:
    """Read CSV, TSV, gzip CSV/TSV, XLSX, or iterable mapping records.

    Sheet indices are zero-based. A workbook containing multiple worksheets
    requires an explicit sheet. The default comma delimiter becomes a tab for
    ``.tsv`` and ``.tsv.gz`` paths. Record hashes identify canonical, type-tagged
    JSON bytes, independently of dictionary key order, and ``size`` is the byte
    length of that representation. Their row order remains significant.
    XLSX worksheet bounding rectangles above 10,000,000 cells are unsupported,
    including formatting-only padding, to bound rectangle materialization.
    Explicit ``format`` takes precedence over filename extensions and accepts
    csv, tsv, xlsx, records, and their table-* aliases.
    """
    if isinstance(source, (str, os.PathLike)):
        return _load_file(
            source,
            format=format,
            sheet=sheet,
            delimiter=delimiter,
            encoding=encoding,
        )
    table = _load_records(source)
    try:
        kind = _normalize_kind(format)
        if kind is not None and kind != "records":
            raise _invalid(
                "In-memory mapping records require the records format.",
                "source_format_mismatch",
                status=ImportStatus.MAPPING_INVALID,
                stage=ImportStage.MAPPING,
            )
        if sheet is not None:
            raise _invalid(
                "Sheet selection applies only to XLSX files.",
                "unexpected_sheet",
                status=ImportStatus.MAPPING_INVALID,
                stage=ImportStage.MAPPING,
            )
    except TableSourceError as error:
        error.sha256 = table.sha256
        error.size = table.size
        raise
    return table


def _normalize_kind(value: str | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        normalized = value.strip().lower().lstrip(".")
        if normalized.startswith("table-"):
            normalized = normalized[len("table-") :]
        if normalized in {"csv", "tsv", "xlsx", "records"}:
            return normalized
    raise _invalid(
        "Format must be csv, tsv, xlsx, records, or a table-* alias.",
        "invalid_format",
        status=ImportStatus.MAPPING_INVALID,
        stage=ImportStage.MAPPING,
    )


def _file_kind(path: Path) -> str | None:
    name = path.name.lower()
    if name.endswith((".csv", ".csv.gz")):
        return "csv"
    if name.endswith((".tsv", ".tsv.gz")):
        return "tsv"
    if name.endswith(".xlsx"):
        return "xlsx"
    return None


def _load_file(
    source: str | os.PathLike[str],
    *,
    format: str | None,
    sheet: str | int | None,
    delimiter: str,
    encoding: str,
) -> TableData:
    try:
        source_text = os.fspath(source)
        if not isinstance(source_text, str) or not source_text.strip():
            raise ValueError("Source must resolve to a nonempty string path.")
        path = Path(source_text)
    except (TypeError, ValueError, OSError) as exc:
        raise TableSourceError(
            str(exc),
            status=ImportStatus.UNAVAILABLE,
            stage=ImportStage.SOURCE,
            code="source_invalid",
            source="invalid:path",
            kind=None,
        ) from exc
    kind = _file_kind(path)
    format_error: TableSourceError | None = None
    try:
        kind = _normalize_kind(format) or kind
    except TableSourceError as exc:
        # A readable file still supplies evidence for an invalid format option.
        format_error = exc
    try:
        raw = path.read_bytes()
    except (OSError, ValueError) as exc:
        raise TableSourceError(
            str(exc),
            status=ImportStatus.UNAVAILABLE,
            stage=ImportStage.SOURCE,
            code="source_unavailable",
            source=source_text,
            kind=kind,
        ) from exc
    digest = hashlib.sha256(raw).hexdigest()
    size = len(raw)
    try:
        if format_error is not None:
            raise format_error
        if kind == "records":
            raise _invalid(
                "File sources require a CSV, TSV, or XLSX format.",
                "source_format_mismatch",
                status=ImportStatus.MAPPING_INVALID,
                stage=ImportStage.MAPPING,
            )
        if kind is None:
            raise _invalid(
                "Expected a CSV, TSV, CSV/TSV gzip, or XLSX file.",
                "unsupported_format",
                status=ImportStatus.UNSUPPORTED,
                stage=ImportStage.SOURCE,
            )
        if kind == "xlsx":
            fields, rows, changes = _load_xlsx(raw, sheet)
        else:
            if sheet is not None:
                raise _invalid(
                    "Sheet selection applies only to XLSX files.",
                    "unexpected_sheet",
                    status=ImportStatus.MAPPING_INVALID,
                    stage=ImportStage.MAPPING,
                )
            changes: tuple[Transformation, ...] = ()
            if path.name.lower().endswith(".gz"):
                try:
                    raw = gzip.decompress(raw)
                except (OSError, EOFError, zlib.error) as exc:
                    raise _invalid(
                        str(exc),
                        "gzip_invalid",
                        status=ImportStatus.SYNTAX_INVALID,
                        stage=ImportStage.SYNTAX,
                    ) from exc
                changes = (
                    Transformation(
                        "gzip_decompressed",
                        "Decompressed the gzip table source.",
                    ),
                )
            if kind == "tsv" and delimiter == ",":
                delimiter = "\t"
            fields, rows, csv_changes = _load_csv(raw, delimiter, encoding)
            changes += csv_changes
    except TableSourceError as exc:
        exc.source = source_text
        exc.kind = kind
        exc.sha256 = digest
        exc.size = size
        raise
    return TableData(fields, rows, source_text, kind, digest, size, changes)


def _headers(values: tuple[object, ...]) -> tuple[str, ...]:
    if not values:
        raise _invalid("The table requires a header row.", "missing_header")
    seen: set[str] = set()
    fields: list[str] = []
    for index, value in enumerate(values):
        at = ("header", str(index))
        if not isinstance(value, str):
            raise _invalid("Headers must be strings.", "invalid_header", at)
        if not value.strip():
            raise _invalid("Headers must not be blank.", "blank_header", at)
        if value in seen:
            raise _invalid("Headers must be unique.", "duplicate_header", at)
        seen.add(value)
        fields.append(value)
    return tuple(fields)


def _check_csv_quotes(text: str, delimiter: str) -> None:
    """Reject malformed quote placement that csv.reader(strict=True) accepts."""
    state = "start"
    line = 1
    previous = ""
    for char in text:
        if state == "quoted":
            if char == '"':
                state = "closed"
        elif state == "closed":
            if char == '"':
                state = "quoted"
            elif char == delimiter or char in "\r\n":
                state = "start"
            else:
                raise _invalid(
                    "Unexpected character after a closing quote.",
                    "csv_invalid",
                    ("line", str(line)),
                    status=ImportStatus.SYNTAX_INVALID,
                    stage=ImportStage.SYNTAX,
                )
        elif char == '"':
            if state != "start":
                raise _invalid(
                    "A quote must begin at the start of a field.",
                    "csv_invalid",
                    ("line", str(line)),
                    status=ImportStatus.SYNTAX_INVALID,
                    stage=ImportStage.SYNTAX,
                )
            state = "quoted"
        elif char == delimiter or char in "\r\n":
            state = "start"
        else:
            state = "plain"
        if char == "\r" or (char == "\n" and previous != "\r"):
            line += 1
        previous = char
    if state == "quoted":
        raise _invalid(
            "Unterminated quoted field.",
            "csv_invalid",
            ("line", str(line)),
            status=ImportStatus.SYNTAX_INVALID,
            stage=ImportStage.SYNTAX,
        )


def _load_csv(
    raw: bytes,
    delimiter: str,
    encoding: str,
) -> tuple[tuple[str, ...], tuple[dict[str, object], ...], tuple[Transformation, ...]]:
    if (
        not isinstance(delimiter, str)
        or len(delimiter) != 1
        or delimiter in '\r\n\x00"'
    ):
        raise _invalid(
            "Delimiter must be one character other than newline, NUL, or quote.",
            "invalid_delimiter",
            status=ImportStatus.MAPPING_INVALID,
            stage=ImportStage.MAPPING,
        )
    try:
        if not isinstance(encoding, str) or not encoding.strip() or "\x00" in encoding:
            raise LookupError("Encoding must be a nonempty codec name without NUL.")
        decoded = raw.decode(encoding)
    except LookupError as exc:
        raise _invalid(
            str(exc),
            "invalid_encoding",
            status=ImportStatus.MAPPING_INVALID,
            stage=ImportStage.MAPPING,
        ) from exc
    except UnicodeError as exc:
        raise _invalid(
            str(exc),
            "text_decode_failed",
            status=ImportStatus.SYNTAX_INVALID,
            stage=ImportStage.SYNTAX,
        ) from exc
    changes: tuple[Transformation, ...] = ()
    if decoded.startswith("\ufeff"):
        decoded = decoded[1:]
        changes = (
            Transformation(
                "text_bom_removed",
                "Removed the leading Unicode byte order mark.",
                count=1,
            ),
        )
    _check_csv_quotes(decoded, delimiter)
    reader = csv.reader(
        io.StringIO(decoded, newline=""), delimiter=delimiter, strict=True
    )
    try:
        fields = _headers(tuple(next(reader, ())))
        rows: list[dict[str, object]] = []
        for row in reader:
            if len(row) != len(fields):
                raise _invalid(
                    f"Expected {len(fields)} cells; found {len(row)}.",
                    "row_width_mismatch",
                    ("line", str(reader.line_num)),
                )
            rows.append(dict(zip(fields, row)))
    except csv.Error as exc:
        raise _invalid(
            str(exc),
            "csv_invalid",
            ("line", str(reader.line_num)),
            status=ImportStatus.SYNTAX_INVALID,
            stage=ImportStage.SYNTAX,
        ) from exc
    return fields, tuple(rows), changes


def _load_xlsx(
    raw: bytes,
    sheet: str | int | None,
) -> tuple[tuple[str, ...], tuple[dict[str, object], ...], tuple[Transformation, ...]]:
    # This is deliberately the only optional dependency import in the module.
    try:
        import openpyxl
    except ImportError as exc:
        raise _invalid(
            "Reading XLSX requires the optional dependency: pip install 'pix[excel]'.",
            "xlsx_dependency_unavailable",
            status=ImportStatus.UNAVAILABLE,
            stage=ImportStage.SOURCE,
        ) from exc
    try:
        _check_xlsx_coordinates(raw, openpyxl.utils.cell.coordinate_to_tuple)
        workbook = openpyxl.load_workbook(
            io.BytesIO(raw),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except TableSourceError:
        raise
    except Exception as exc:
        # openpyxl delegates ZIP, XML, and cell parsing to several libraries.
        raise _invalid(
            str(exc),
            "xlsx_invalid",
            status=ImportStatus.SYNTAX_INVALID,
            stage=ImportStage.SYNTAX,
        ) from exc
    try:
        worksheets = workbook.worksheets
        if sheet is None:
            if len(worksheets) != 1:
                raise _invalid(
                    "Choose a sheet explicitly when the workbook does not have "
                    "exactly one worksheet.",
                    "sheet_required",
                    status=ImportStatus.MAPPING_INVALID,
                    stage=ImportStage.MAPPING,
                )
            worksheet = worksheets[0]
        elif isinstance(sheet, str):
            matches = [value for value in worksheets if value.title == sheet]
            if not matches:
                raise _invalid(
                    "The requested worksheet does not exist.",
                    "sheet_not_found",
                    status=ImportStatus.MAPPING_INVALID,
                    stage=ImportStage.MAPPING,
                )
            worksheet = matches[0]
        elif isinstance(sheet, int) and not isinstance(sheet, bool):
            if not 0 <= sheet < len(worksheets):
                raise _invalid(
                    "The zero-based worksheet index is out of range.",
                    "sheet_not_found",
                    status=ImportStatus.MAPPING_INVALID,
                    stage=ImportStage.MAPPING,
                )
            worksheet = worksheets[sheet]
        else:
            raise _invalid(
                "Sheet must be a worksheet name or a zero-based integer index.",
                "invalid_sheet",
                status=ImportStatus.MAPPING_INVALID,
                stage=ImportStage.MAPPING,
            )
        if worksheet.merged_cells.ranges:
            raise _invalid(
                "Merged worksheet cells cannot be interpreted as a table.",
                "xlsx_merged_cells",
                ("sheet", worksheet.title),
                status=ImportStatus.UNSUPPORTED,
                stage=ImportStage.SOURCE,
            )
        if worksheet.max_row * worksheet.max_column > _MAX_XLSX_CELLS:
            raise _invalid(
                "Worksheet dimensions exceed the supported bounding rectangle "
                f"of {_MAX_XLSX_CELLS:,} cells, including formatting-only padding.",
                "xlsx_dimensions_exceeded",
                ("sheet", worksheet.title),
                status=ImportStatus.UNSUPPORTED,
                stage=ImportStage.SOURCE,
            )
        values: dict[int, dict[int, object]] = {}
        width = 0
        cells = (cell for row in worksheet.iter_rows() for cell in row)
        for cell in cells:
            at = ("sheet", worksheet.title, "cell", cell.coordinate)
            if cell.data_type == "f":
                raise _invalid(
                    "Formula cells require an explicit value-only source.",
                    "xlsx_formula",
                    at,
                    status=ImportStatus.UNSUPPORTED,
                    stage=ImportStage.SOURCE,
                )
            if cell.data_type == "e":
                raise _invalid(
                    "Worksheet contains an error cell.", "xlsx_cell_error", at
                )
            value = cell.value
            # Empty inline strings are explicit cells, unlike absent cells or
            # formatting-only padding. Keep them visible to header validation.
            if value is None and cell.data_type in {"inlineStr", "s"}:
                value = ""
            if value is not None:
                width = max(width, cell.column)
                values.setdefault(cell.row, {})[cell.column] = value
        if not values:
            raise _invalid("The table requires a header row.", "missing_header")
        header = values.get(1, {})
        fields = _headers(tuple(header.get(column) for column in range(1, width + 1)))
        records: list[dict[str, object]] = []
        for expected_index, index in enumerate(sorted(values)[1:], start=2):
            row = values[index]
            if index != expected_index or all(value == "" for value in row.values()):
                raise _invalid(
                    "An empty row inside the worksheet table is ambiguous.",
                    "xlsx_empty_row",
                    ("sheet", worksheet.title, "row", str(expected_index)),
                )
            record: dict[str, object] = {}
            for column, field in enumerate(fields, start=1):
                value = row.get(column)
                snapshot, _ = _cell(value, ("rows", str(index - 2), field), set())
                record[field] = snapshot
            records.append(record)
        changes = (
            Transformation(
                "xlsx_sheet_selected",
                f"Read worksheet {worksheet.title!r}.",
                at=("sheet", worksheet.title),
                count=1,
            ),
        )
        return fields, tuple(records), changes
    finally:
        workbook.close()


def _check_xlsx_coordinates(raw: bytes, coordinate_to_tuple: Any) -> None:
    """Reject duplicate coordinates before openpyxl can overwrite their cells."""
    content_ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise _invalid(
                "XLSX archive contains duplicate part names.",
                "xlsx_duplicate_archive_part",
                status=ImportStatus.SYNTAX_INVALID,
                stage=ImportStage.SYNTAX,
            )
        content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        overrides = {
            value.attrib["PartName"].lstrip("/"): value.attrib["ContentType"]
            for value in content_types.findall(f"{content_ns}Override")
        }
        defaults = {
            value.attrib["Extension"]: value.attrib["ContentType"]
            for value in content_types.findall(f"{content_ns}Default")
        }
        worksheet_names = {
            name
            for name in names
            if overrides.get(name, defaults.get(name.rsplit(".", 1)[-1], "")).endswith(
                ".worksheet+xml"
            )
        }
        # Workbook relationships also identify worksheet parts independently
        # of content-type declarations (which third-party producers can vary).
        relation_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        for name in names:
            if not name.endswith(".rels"):
                continue
            relationships = ElementTree.fromstring(archive.read(name))
            for relation in relationships.findall(f"{relation_ns}Relationship"):
                if not relation.attrib.get("Type", "").endswith("/worksheet"):
                    continue
                if relation.attrib.get("TargetMode") == "External":
                    continue
                target = relation.attrib["Target"]
                if target.startswith("/"):
                    worksheet_names.add(target.lstrip("/"))
                else:
                    worksheet_names.add(
                        posixpath.normpath(
                            posixpath.join(
                                posixpath.dirname(posixpath.dirname(name)),
                                target,
                            )
                        )
                    )
        for name in sorted(worksheet_names):
            root = ElementTree.fromstring(archive.read(name))
            namespace = root.tag.rsplit("}", 1)[0] + "}" if "}" in root.tag else ""
            sheet_data = root.find(f"{namespace}sheetData")
            if sheet_data is None:
                continue
            seen_rows: set[int] = set()
            seen_cells: set[tuple[int, int]] = set()
            row_index = 0
            for row in sheet_data.findall(f"{namespace}row"):
                row_index = int(row.attrib.get("r", row_index + 1))
                if row_index in seen_rows:
                    raise _invalid(
                        "Worksheet XML contains a duplicate row coordinate.",
                        "xlsx_duplicate_row",
                        ("part", name, "row", str(row_index)),
                        status=ImportStatus.SYNTAX_INVALID,
                        stage=ImportStage.SYNTAX,
                    )
                seen_rows.add(row_index)
                column_index = 0
                for cell in row.findall(f"{namespace}c"):
                    reference = cell.attrib.get("r")
                    if reference is None:
                        column_index += 1
                        coordinate = (row_index, column_index)
                    else:
                        coordinate = coordinate_to_tuple(reference)
                        column_index = coordinate[1]
                    if coordinate in seen_cells:
                        raise _invalid(
                            "Worksheet XML contains a duplicate cell coordinate.",
                            "xlsx_duplicate_cell",
                            ("part", name, "cell", reference or str(coordinate)),
                            status=ImportStatus.SYNTAX_INVALID,
                            stage=ImportStage.SYNTAX,
                        )
                    seen_cells.add(coordinate)


def _cell(
    value: object,
    at: tuple[str, ...],
    active: set[int],
) -> tuple[object, object]:
    """Snapshot supported native values and give each type an explicit tag."""
    if value is None:
        return value, ["null"]
    if isinstance(value, bool):
        return value, ["bool", value]
    if isinstance(value, str):
        return value, ["str", value]
    if isinstance(value, int):
        return value, ["int", str(value)]
    if isinstance(value, float) and math.isfinite(value):
        return value, ["float", value.hex()]
    if isinstance(value, datetime):
        return value, ["datetime", value.isoformat(), value.fold]
    if isinstance(value, date):
        return value, ["date", value.isoformat()]
    if isinstance(value, Decimal) and value.is_finite():
        return value, ["decimal", str(value)]
    if isinstance(value, (list, tuple)):
        if id(value) in active:
            raise _invalid(
                "Record cell contains a cyclic sequence.", "unsupported_cell", at
            )
        active.add(id(value))
        try:
            snapshot: list[object] = []
            tagged: list[object] = []
            for index, item in enumerate(value):
                copied, encoded = _cell(item, (*at, str(index)), active)
                snapshot.append(copied)
                tagged.append(encoded)
        finally:
            active.remove(id(value))
        if isinstance(value, tuple):
            return tuple(snapshot), ["tuple", tagged]
        return snapshot, ["list", tagged]
    raise _invalid(
        "Cell values must be finite numbers, strings, booleans, dates, "
        "datetimes, null, or lists/tuples of these values.",
        "unsupported_cell",
        at,
    )


def _load_records(source: object) -> TableData:
    if isinstance(source, (Mapping, bytes, bytearray)):
        raise _invalid("Expected an iterable of mapping records.", "invalid_records")
    try:
        iterator = iter(source)  # type: ignore[arg-type]
    except TypeError as exc:
        raise _invalid(
            "Expected an iterable of mapping records.", "invalid_records"
        ) from exc
    except Exception as exc:
        raise _invalid(
            str(exc),
            "records_unreadable",
            status=ImportStatus.UNAVAILABLE,
            stage=ImportStage.SOURCE,
        ) from exc
    fields: dict[str, None] = {}
    rows: list[dict[str, object]] = []
    encoded_rows: list[Any] = []
    try:
        for index, value in enumerate(iterator):
            at = ("rows", str(index))
            if not isinstance(value, Mapping):
                raise _invalid("Every record must be a mapping.", "invalid_record", at)
            record: dict[str, object] = {}
            tags: dict[str, object] = {}
            for field, cell in value.items():
                if not isinstance(field, str) or not field.strip():
                    raise _invalid(
                        "Record field names must be nonblank strings.",
                        "invalid_record_field",
                        at,
                    )
                copied, encoded = _cell(cell, (*at, field), set())
                record[field] = copied
                tags[field] = encoded
                fields.setdefault(field, None)
            rows.append(record)
            encoded_rows.append([[field, tags[field]] for field in sorted(tags)])
    except TableSourceError:
        raise
    except Exception as exc:
        raise _invalid(
            str(exc),
            "records_unreadable",
            ("rows", str(len(rows))),
            status=ImportStatus.UNAVAILABLE,
            stage=ImportStage.SOURCE,
        ) from exc
    raw = json.dumps(
        ["pix-table-records-v1", encoded_rows],
        ensure_ascii=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("ascii")
    return TableData(
        tuple(fields),
        tuple(rows),
        "memory:records",
        "records",
        hashlib.sha256(raw).hexdigest(),
        len(raw),
    )


__all__ = ["TableData", "TableSourceError", "load_table"]
